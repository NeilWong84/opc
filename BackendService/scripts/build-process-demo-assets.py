"""Create inspectable synthetic attachment bytes; run with the bundled artifact Python.

Office spreadsheets/slides are authored by build-process-demo-office.mjs afterwards.
All output is local and ignored. No live accounts or customer material are read.
"""
from hashlib import sha256
import json
from pathlib import Path
import subprocess
from urllib.request import Request, urlopen
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1] / '.local/process-demo-20261025'
ASSETS = ROOT / 'assets'
QA = ROOT / 'qa'
FAMILIES = [
    {'key': 'booking', 'title': '店舗予約管理ツール', 'zh': '门店预约管理工具',
     'goal': '予約の二重登録と手作業の転記を減らす。减少重复预约和人工抄录。',
     'features': ['予約一覧と日付検索 / 预约列表与日期筛选', '同一時間帯の重複防止 / 阻止同一时段重复预约', 'CSV 出力と再読込 / CSV 导出与重新读取'],
     'excluded': 'オンライン決済、実在する顧客情報、本番公開は含まない。',
     'rows': [['B001', '2026-11-02', '10:00', 'Demo A', 'confirmed'], ['B002', '2026-11-02', '11:00', 'Demo B', 'pending'], ['B003', '2026-11-03', '14:00', 'Demo C', 'confirmed']]},
    {'key': 'media', 'title': '社内メディア資料ポータル', 'zh': '内部多媒体资料门户',
     'goal': '画像、音声、動画と説明資料を権限別に整理する。按权限整理图片、音频、视频和说明材料。',
     'features': ['形式別一覧と検索 / 按格式筛选与检索', '権限確認後の原本取得 / 鉴权后下载原件', '画像と MP4 のプレビュー / 图片与 MP4 预览'],
     'excluded': '音声文字起こし、Office 変換、OSS 接続、本番公開は含まない。',
     'rows': [['M001', 'image', 'landscape.jpg', 'reference', 'clean'], ['M002', 'audio', 'sample-tone.mp3', 'reference', 'clean'], ['M003', 'video', 'sample-road.mp4', 'reference', 'clean']]},
]


def document(family):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Inches(.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        if style.element.pPr is not None:
            for border in list(style.element.pPr.findall(qn('w:pBdr'))):
                style.element.pPr.remove(border)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.space_after = Pt(7)
    doc.add_paragraph(family['title'], 'Title')
    doc.add_paragraph(family['zh'] + '  要件説明書  2026年10月25日')
    doc.add_paragraph('合成デモ資料 / 合成演示资料。本文用于 OPC 需求评审及附件展示，不含真实客户资料。')
    doc.add_heading('目的と範囲', 1)
    doc.add_paragraph(family['goal'])
    for text in family['features']:
        doc.add_paragraph(text, 'List Bullet')
    doc.add_heading('検収項目', 1)
    table = doc.add_table(rows=1, cols=3)
    table.columns[0].width, table.columns[1].width, table.columns[2].width = Inches(.6), Inches(3.2), Inches(2.2)
    for cell, text in zip(table.rows[0].cells, ['ID', '確認内容 / 检查内容', '期待結果 / 预期结果']):
        cell.text = text
    for row in [('A01', '保存後に再取得 / 保存后重新读取', '同一内容を取得'), ('A02', '別アカウントで私有資料へアクセス', 'アクセスを拒否'), ('A03', '添付原本をダウンロード', 'SHA-256 が一致')]:
        for cell, text in zip(table.add_row().cells, row):
            cell.text = text
    for i, row in enumerate(table.rows):
        for cell in row.cells:
            tcpr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd')
            shade.set(qn('w:fill'), 'DCE6F1' if i == 0 else ('F3F6FA' if i % 2 else 'FFFFFF'))
            tcpr.append(shade)
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                item = OxmlElement('w:' + edge)
                for k, v in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
                    item.set(qn('w:' + k), v)
                borders.append(item)
            tcpr.append(borders)
    doc.add_heading('対象外と資料の扱い', 1)
    doc.add_paragraph(family['excluded'])
    doc.add_paragraph('客户私有附件只供受理使用。经审核明确选中的资料才进入大厅；交付经内部质检与对客发布后供客户验收。')
    doc.add_heading('補足と次の確認', 1)
    doc.add_paragraph('利用人数は 3 名、データは合成値。范围变更重新整理审核；JPY 参考报价不表示实际收款。')
    doc.save(ASSETS / (family['key'] + '-requirements.docx'))


def pdf(family):
    pdfmetrics.registerFont(TTFont('DemoCJK', 'C:/Windows/Fonts/msyh.ttc', subfontIndex=0))
    path = ASSETS / (family['key'] + '-acceptance.pdf')
    c = canvas.Canvas(str(path), pagesize=(595, 842))
    c.setTitle(family['title'] + ' 検収ガイド')
    c.setFont('DemoCJK', 22)
    c.drawString(48, 776, family['title'])
    c.setFont('DemoCJK', 12)
    for y, line in zip(range(730, 180, -32), [
        '検収ガイド / 验收指南', '2026-10-25   合成デモ資料 / 合成演示资料',
        '1  合意した範囲と最新納品版を確認する',
        '   确认约定范围与最新交付版本。',
        '2  サンプル値を入力し保存後に再取得する',
        '   使用合成值验证保存、筛选与重新读取。',
        '3  権限のない資料へアクセスできないことを確認する',
        '   验证私有附件和内部会话不可越权访问。',
        '4  原本を取得し SHA-256 と照合する',
        '   下载原件并比对清单中的 SHA-256。',
        '5  修正箇所または検収の意思を明確に記録する',
        '   记录具体修改意见或明确验收，不触发付款。',
        '注意  本番データ、個人情報、実決済は使用しない。',
        'Office / PDF のアプリ内変換と OSS 接続は対象外。',
    ]):
        c.drawString(48, y, line)
    c.setFont('DemoCJK', 10)
    c.drawString(48, 45, 'OPC  Synthetic development fixture                                      1 / 1')
    c.save()


def screenshot(family):
    image = Image.new('RGB', (1280, 800), '#f3f6fc')
    draw = ImageDraw.Draw(image)
    def font(size):
        return ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', size)
    draw.rectangle((0, 0, 1280, 90), fill='#173455')
    draw.text((40, 24), family['title'], fill='white', font=font(30))
    draw.text((40, 112), '画面参考 / 页面示意   合成デモ資料', fill='#3d526c', font=font(23))
    draw.rounded_rectangle((40, 180, 1240, 710), radius=14, fill='white')
    headers = ['ID', '日付 / 種類', '時間 / ファイル名', '表示名 / 範囲', '状態']
    xs = [65, 235, 465, 835, 1085]
    for x, text in zip(xs, headers):
        draw.text((x, 215), text, fill='#1c3559', font=font(20))
    for i, row in enumerate(family['rows']):
        y = 300 + i * 100
        draw.line((65, y - 20, 1210, y - 20), fill='#dce3ee', width=2)
        for x, text in zip(xs, row):
            draw.text((x, y), text, fill='#334155', font=font(19))
    draw.text((40, 750), '架空データのみ / 仅合成数据     UI reference for development review', fill='#526279', font=font(20))
    image.save(ASSETS / (family['key'] + '-screen.png'))


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    QA.mkdir(exist_ok=True)
    (ROOT / 'families.json').write_text(json.dumps(FAMILIES, ensure_ascii=False, indent=2), encoding='utf-8')
    for family in FAMILIES:
        document(family)
        pdf(family)
        screenshot(family)
        import csv
        with (ASSETS / (family['key'] + '-sample.csv')).open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(['id', 'date_or_kind', 'time_or_filename', 'name_or_access', 'status'])
            writer.writerows(family['rows'])
    sources = []
    for filename, url, page in [
        ('landscape.jpg', 'https://samplelib.com/jpeg/sample-photo-1280x720.jpg', 'https://samplelib.com/sample-jpeg.html'),
        ('sample-tone.mp3', 'https://samplelib.com/mp3/sample-3s.mp3', 'https://samplelib.com/sample-mp3.html'),
        ('sample-road.mp4', 'https://samplelib.com/mp4/sample-5s-360p.mp4', 'https://samplelib.com/sample-mp4.html'),
    ]:
        target = ASSETS / filename
        if not target.exists():
            with urlopen(Request(url, headers={'User-Agent': 'OPC-Local-Demo/1.0'}), timeout=60) as response:
                content = response.read(25 * 1024 * 1024 + 1)
            if not content or len(content) > 25 * 1024 * 1024:
                raise ValueError('Download size outside fixture limit')
            target.write_bytes(content)
        sources.append({'file': filename, 'url': url, 'source_page': page, 'license_url': 'https://samplelib.com/license.html', 'retrieved_on': '2026-10-25', 'sha256': sha256(target.read_bytes()).hexdigest(), 'size_bytes': target.stat().st_size})
    (ASSETS / 'media-sources.md').write_text('# 演示素材来源\n\nSamplelib 原创测试素材，许可页声明可自由使用。\n图片为风景参考，音频为合成旋律，视频为道路测试片段，不是真实客户录音或项目操作录屏。\n\n' + '\n\n'.join(f"- {s['file']}\n  来源：{s['source_page']}\n  原件：{s['url']}\n  许可：{s['license_url']}\n  SHA-256：{s['sha256']}" for s in sources), encoding='utf-8')
    (ASSETS / 'private-intake-note.txt').write_text('客户私有受理补充 / お客様の非公開補足\n合成数据，无个人信息。利用人数3名；预算仅参考报价；此文件不得随大厅发布。\n', encoding='utf-8')
    (ASSETS / 'media-test-notes.txt').write_text('メディア検証メモ / 多媒体验证说明\nMP3 は3秒の合成旋律。MP4 は道路のテスト素材。画像は風景。\nこれらは形式と原本取得の検証用であり、顧客の発言や実装の実行証拠ではありません。\n', encoding='utf-8')
    source = '''"""Synthetic slot store. No network, credentials or real customer data."""\nimport json\nfrom pathlib import Path\n\ndef reserve(path, slot):\n    path = Path(path)\n    rows = json.loads(path.read_text()) if path.exists() else []\n    if slot in rows:\n        raise ValueError("duplicate slot")\n    rows.append(slot)\n    path.write_text(json.dumps(rows), encoding="utf-8")\n    return rows\n'''
    tests = '''import tempfile\nimport unittest\nfrom pathlib import Path\nfrom booking import reserve\n\nclass BookingTest(unittest.TestCase):\n    def test_save_reload_and_duplicate(self):\n        with tempfile.TemporaryDirectory() as root:\n            path = Path(root) / "slots.json"\n            self.assertEqual(reserve(path, "2026-11-02T10:00"), ["2026-11-02T10:00"])\n            with self.assertRaises(ValueError):\n                reserve(path, "2026-11-02T10:00")\n            self.assertEqual(len(reserve(path, "2026-11-02T11:00")), 2)\n\nif __name__ == "__main__":\n    unittest.main()\n'''
    check = ROOT / 'source-check'
    check.mkdir(exist_ok=True)
    (check / 'booking.py').write_text(source, encoding='utf-8')
    (check / 'test_booking.py').write_text(tests, encoding='utf-8')
    import sys
    result = subprocess.run([sys.executable, '-m', 'unittest', '-v'], cwd=check, text=True, capture_output=True, check=True)
    (QA / 'source-tests.txt').write_text(result.stdout + result.stderr, encoding='utf-8')
    for version in [1, 2]:
        with ZipFile(ASSETS / f'booking-source-v{version}.zip', 'w', ZIP_DEFLATED) as archive:
            archive.writestr('booking.py', source)
            archive.writestr('test_booking.py', tests)
            archive.writestr('README.md', f'# Synthetic booking fixture v{version}\nPython 3 standard library only.\nRun: python -m unittest -v\nThis is a minimal persistence example, not a production service.\n')
            archive.writestr('CHANGELOG.md', 'v2: clarify duplicate-slot validation in documentation.\n' if version == 2 else 'v1: synthetic persistence example.\n')
    (ROOT / 'sources.json').write_text(json.dumps(sources, ensure_ascii=False, indent=2), encoding='utf-8')
    print('Base assets generated; source unit test passed. Run Office builder then inspection before seeding.')


if __name__ == '__main__':
    main()
