"""Rebuild original, non-customer sample attachments using the bundled runtime.

Requires python-docx, reportlab, Pillow, the supplied @oai/artifact-tool Node
runtime and ffmpeg. Temporary render/validation files stay outside the repo.
"""
from __future__ import annotations

import csv
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml.ns import qn
from PIL import Image
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "APPClient/assets/references"
SOURCE = ROOT / "Mockup/task-reference-assets"
RUNTIME = Path(os.environ.get("OPC_ARTIFACT_RUNTIME", Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies"))
NODE = RUNTIME / "node/bin/node.exe"
SKILLS = Path.home() / ".codex/plugins/cache/openai-primary-runtime"
BUILD = Path(tempfile.mkdtemp(prefix="opc-reference-samples-"))


NODE_BUILDER = r'''
import fs from 'node:fs/promises';
import path from 'node:path';
import {pathToFileURL} from 'node:url';
import {Workbook, SpreadsheetFile, Presentation, PresentationFile, FileBlob} from '@oai/artifact-tool';
const [out, build, skill, python] = process.argv.slice(2);
const fields = [['項目名','型','必須'],['問い合わせID','文字列','はい'],['件名','文字列','はい'],['本文','長文','はい'],['ステータス','選択','はい'],['担当者','ユーザー','いいえ'],['受付日時','日時','自動'],['更新日時','日時','自動']];
const statuses = [['状態値','表示名','内容'],['new','新規','問い合わせの受付直後'],['in_progress','対応中','担当者が内容を確認'],['resolved','完了','対応履歴を記録済み']];
const permissions = [['操作','担当者','管理担当'],['問い合わせ閲覧','可','可'],['担当者変更','可','可'],['チーム設定','不可','可']];
function style(sheet, rows, cols) {
  sheet.showGridLines=false;
  const end=String.fromCharCode(64+cols);
  const all=sheet.getRange(`A1:${end}${rows}`);
  all.format.font={name:'Yu Gothic',size:11,color:'#183E35'};
  all.format.rowHeight=28;
  all.format.columnWidth=20;
  all.format.verticalAlignment='center';
  sheet.getRange(`A1:${end}1`).format.fill='#086E61';
  sheet.getRange(`A1:${end}1`).format.font={name:'Yu Gothic',size:11,bold:true,color:'#FFFFFF'};
}
async function exportBook(name, sheets) {
  const wb=Workbook.create();
  for(const [sheetName, values] of sheets) {
    const s=wb.worksheets.add(sheetName);
    s.getRange('A1').write(values);
    style(s,values.length,values[0].length);
    s.getRange('A1:A20').format.columnWidth=25;
    if(values[0].length===4) s.getRange('D1:D20').format.columnWidth=49;
  }
  if(name==='quote-template.xlsx') {
    const s=wb.worksheets.getItemAt(0);
    s.getRange('D2:D3').formulas=[['=B2*C2'],['=B3*C3']];
    s.getRange('D4').formulas=[['=SUM(D2:D3)']];
    s.getRange('B2:C3').format.font={color:'#1664B7',name:'Yu Gothic',size:11};
    s.getRange('C2:D4').setNumberFormat('#,##0');
    s.getRange('D1:D20').format.columnWidth=24;
  }
  wb.recalculate();
  console.log((await wb.inspect({kind:'table',range:`'${sheets[0][0]}'!A1:D9`,include:'values,formulas',maxChars:1800})).ndjson);
  console.log((await wb.inspect({kind:'match',searchTerm:'#REF!|#DIV/0!|#VALUE!|#NAME\\?',options:{useRegex:true,maxResults:20},maxChars:600})).ndjson);
  for(const [sheetName,values] of sheets) {
    const blob=await wb.render({sheetName,range:`A1:${String.fromCharCode(64+values[0].length)}${values.length}`,scale:1.5,format:'png'});
    await fs.writeFile(path.join(build,`${name}-${sheetName}.png`),new Uint8Array(await blob.arrayBuffer()));
  }
  await (await SpreadsheetFile.exportXlsx(wb)).save(path.join(build,name));
  await fs.copyFile(path.join(build,name),path.join(out,name));
}
await exportBook('fields.xlsx',[['項目定義',fields],['ステータス',statuses],['権限',permissions]]);
await exportBook('quote-template.xlsx',[['見積明細',[['品名','数量','単価（円）','金額（円）'],['要件整理',1,30000,null],['開発作業',1,120000,null],['合計',null,null,null]]]]);
const presentation=Presentation.create({slideSize:{width:1280,height:800}});
function text(slide,value,left,top,width,height,size,color='#183E35') {
  const shape=slide.shapes.add({geometry:'textbox',position:{left,top,width,height},fill:'none',line:{fill:'none',width:0}});
  shape.text=value;
  shape.text.style={typeface:'Yu Gothic',fontSize:size,color,autoFit:'none'};
}
const cover=presentation.slides.add(); cover.background.fill='#183E35';
text(cover,'INQUIRY WORKSPACE',80,180,1100,100,60,'#FFFFFF');
text(cover,'チームの対応を、ひとつの画面に。',80,320,1100,70,36,'#D7ECE3');
text(cover,'登録     共有     対応',80,460,1100,60,32,'#FFFFFF');
text(cover,'1 / 3    OPC 架空の参考資料',80,730,1100,30,18,'#D7ECE3');
for(const [i,name,title,note] of [[1,'inquiry-overview.png','02 サービスの全体像','問い合わせ一覧から、担当者と対応状況を確認します。'],[2,'inquiry-detail.png','03 問い合わせの詳細','担当者の変更と対応履歴を記録し、チームに共有します。']]) {
  const slide=presentation.slides.add();slide.background.fill='#F5F7F5';
  text(slide,title,48,30,1184,50,32);
  slide.images.add({blob:new Uint8Array(await fs.readFile(path.join(out,name))),contentType:'image/png',alt:title,fit:'contain',position:{left:120,top:100,width:1040,height:650}});
  text(slide,`${i+1} / 3    ${note}   架空の画面例です。`,48,758,1184,28,17,'#526D63');
  slide.speakerNotes.textFrame.setText('OPC original sample UI. Source: Mockup/task-reference-assets/'+name+'. No real customer data.');
}
const candidatePath=path.join(build,'candidate.pptx');
await (await PresentationFile.exportPptx(presentation)).save(candidatePath);
const {finalizePresentation}=await import(pathToFileURL(path.join(skill,'container_tools/artifact_tool_utils.mjs')).href);
const finalPath=path.join(build,'final','screens.pptx');
await fs.mkdir(path.dirname(finalPath),{recursive:true});
await finalizePresentation({workspaceDir:build,candidatePath,finalPath,pythonExecutable:python,integrityValidatorPath:path.join(skill,'container_tools/inspect_presentation_package_integrity.py'),layoutValidatorPath:path.join(skill,'container_tools/inspect_presentation_layout_geometry.py'),layoutArgs:['--expected-slide-size-emu','12192000,7620000','--validate-heading-fit'],fontPolicy:{basis:'design',families:['Yu Gothic']},explicitTotalSlideCount:3,verifyArtifactToolImport:true,receiptPath:path.join(build,'screens-validation.json')});
await fs.copyFile(finalPath,path.join(out,'screens.pptx'));
const finalDeck=await PresentationFile.importPptx(await FileBlob.load(finalPath));
for(let i=0;i<3;i++) {
  const preview=await finalDeck.export({slide:finalDeck.slides.getItem(i),format:'png',scale:1});
  await fs.writeFile(path.join(build,`slide-${i+1}.png`),new Uint8Array(await preview.arrayBuffer()));
}
'''


def generate():
    OUT.mkdir(parents=True, exist_ok=True)
    for name in ("inquiry-overview.png", "inquiry-detail.png", "workflow-poster.png"):
        shutil.copyfile(SOURCE / name, OUT / name)
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(.75)
    normal = document.styles["Normal"]
    normal.font.name = "Yu Gothic"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    for style_name in ("Title", "Heading 1", "Heading 2"):
        document.styles[style_name].font.name = "Yu Gothic"
        document.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
        document.styles[style_name].element.rPr.rFonts.set(qn("w:eastAsia"), "Yu Gothic")
    pages = [
        ("お問い合わせ管理ツール 要件整理書", ["小規模チームの問い合わせを一元管理し、受付から対応完了までの状況を共有します。", "利用者はチーム内の担当者です。顧客から寄せられた問い合わせを登録し、対応漏れを防ぎます。"], []),
        ("02 対象範囲", ["主な機能", "対象外：外部 CRM 連携・既存データの移行"], ["問い合わせの登録・一覧・検索", "担当者と対応状況の更新", "チーム用ログインと閲覧権限"]),
        ("03 確認ポイント", ["登録から対応完了まで、チームで状況を確認できることを確かめます。"], ["件名と本文から検索できること", "担当者の変更履歴が残ること", "スマートフォンから閲覧できること"]),
    ]
    for index, (title, paragraphs, bullets) in enumerate(pages):
        if index:
            document.add_page_break()
        document.add_paragraph(title, "Title" if index == 0 else "Heading 1")
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        for bullet in bullets:
            document.add_paragraph(bullet, "List Bullet")
        document.add_paragraph(f"OPC 架空の参考資料 v1.0    {index + 1} / 3")
    document.save(OUT / "requirements.docx")

    fixture = (ROOT / "APPClient/lib/data/reference_samples.dart").read_text(encoding="utf-8")
    notes = fixture.split("const referenceNotes = '''", 1)[1].split("''';", 1)[0]
    (OUT / "notes.md").write_text(notes, encoding="utf-8")
    (OUT / "notes.txt").write_text(notes.replace("# ", "").replace("## ", "").replace("`", ""), encoding="utf-8")
    with (OUT / "sample.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(("問い合わせID", "件名", "ステータス", "担当者"))
        writer.writerows((("INQ-001", "資料の送付について", "new", "田中"), ("INQ-002", "利用方法の確認", "in_progress", "佐藤"), ("INQ-003", "登録情報の変更", "resolved", "田中")))
    pdfmetrics.registerFont(TTFont("OPCJapanese", "C:/Windows/Fonts/YuGothR.ttc", subfontIndex=0))
    pdf = canvas.Canvas(str(OUT / "workflow.pdf"), pagesize=(595.28, 841.89))
    pdf.setTitle("お問い合わせ管理ツール 操作の流れ")
    pdf.setAuthor("OPC sample")
    pdf_pages = [
        ("業務フロー 受付から対応完了まで", ["小規模チームの問い合わせを一元管理し、", "受付から対応完了までの状況を共有します。"]),
        ("02 対応の流れ", ["対象外：外部 CRM 連携・既存データの移行", "1. 問い合わせを登録", "2. 担当者が内容と優先度を確認", "3. 対応履歴を記録して完了"]),
        ("03 状態の共有", ["new → in_progress → resolved", "担当者と更新日時を一覧で確認します。"]),
    ]
    for i, (title, paragraphs) in enumerate(pdf_pages):
        pdf.setFillColorRGB(.08, .18, .15)
        pdf.setFont("OPCJapanese", 20)
        pdf.drawString(40, 790, title)
        pdf.setFont("OPCJapanese", 14)
        for row, paragraph in enumerate(paragraphs):
            pdf.drawString(40, 720 - 34 * row, paragraph)
        pdf.setFont("OPCJapanese", 11)
        pdf.drawString(40, 150, "この資料は架空のサンプルです。顧客情報は含みません。")
        pdf.setFont("OPCJapanese", 10)
        pdf.drawString(40, 40, f"OPC 参考資料 v1.0    {i + 1} / 3")
        pdf.showPage()
    pdf.save()

    modules = RUNTIME / "node/node_modules"
    subprocess.run(["cmd", "/c", "mklink", "/J", str(BUILD / "node_modules"), str(modules)], check=True, capture_output=True)
    builder = BUILD / "build.mjs"
    builder.write_text(NODE_BUILDER, encoding="utf-8")
    environment = {**os.environ, "RUNTIME_NODE_MODULES": str(modules)}
    subprocess.run([str(NODE), str(builder), str(OUT), str(BUILD), str(SKILLS / "presentations/26.909.12148/skills/presentations"), str(RUNTIME / "python/python.exe")], check=True, env=environment)

    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-loop", "1", "-t", "4", "-i", str(OUT / "inquiry-overview.png"), "-loop", "1", "-t", "4", "-i", str(OUT / "inquiry-detail.png"), "-filter_complex", "[0:v][1:v]concat=n=2:v=1:a=0,format=yuv420p[v]", "-map", "[v]", "-r", "24", "-t", "8", "-c:v", "libx264", "-preset", "medium", "-crf", "22", "-movflags", "+faststart", str(OUT / "workflow.mp4")], check=True)
    probe = json.loads(subprocess.check_output(["ffprobe", "-v", "quiet", "-show_format", "-show_streams", "-of", "json", str(OUT / "workflow.mp4")], text=True))
    metadata = {"requirements.docx": {"authoredPages": 3}, "fields.xlsx": {"sheets": [{"name": "項目定義", "rows": 8, "columns": 3}, {"name": "ステータス", "rows": 4, "columns": 3}, {"name": "権限", "rows": 4, "columns": 3}]}, "quote-template.xlsx": {"sheets": [{"name": "見積明細", "rows": 4, "columns": 4}], "totalJpy": 150000}, "sample.csv": {"rowsIncludingHeader": 4, "columns": 4}, "screens.pptx": {"slides": 3}, "workflow.pdf": {"pages": 3}, "workflow.mp4": {"durationSeconds": float(probe["format"]["duration"]), "width": probe["streams"][0]["width"], "height": probe["streams"][0]["height"], "codec": probe["streams"][0]["codec_name"]}}
    manifest = {"description": "Original synthetic sample files; no real customer data.", "assets": []}
    for file in sorted(OUT.iterdir()):
        if file.name == "manifest.json" or file.suffix == ".ndjson" or not file.is_file():
            continue
        item = {"file": file.name, "bytes": file.stat().st_size, "sha256": hashlib.sha256(file.read_bytes()).hexdigest(), **metadata.get(file.name, {})}
        if file.suffix == ".png":
            with Image.open(file) as image:
                item.update(width=image.width, height=image.height)
        manifest["assets"].append(item)
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Verification previews: {BUILD}")
    print(f"Manifest: {OUT / 'manifest.json'}")


if __name__ == "__main__":
    generate()

